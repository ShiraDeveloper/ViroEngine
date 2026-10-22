"""Excel exporter - colour-coded, right-to-left workbook (PLAN §9.7). Built in memory only."""

from __future__ import annotations

from datetime import datetime
import io

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from backend.adapters.export.table import CENTRALS_KEY, COLUMNS, HEADERS, Cell, FlatRow, flatten
from backend.core.models import CommitmentRow, ConfidenceLevel, FieldKey, RunReport

SHEET_MAIN = "התחייבויות"
SHEET_REVIEW = "לבדיקה"
SHEET_AUDIT = "יומן שינויים"
SHEET_SUMMARY = "סיכום"

# Same hex values as frontend/src/design/tokens.ts (state.verified / single / review, primary.100).
FILL_VERIFIED = "E3F4E5"
FILL_SINGLE = "F6FBF4"
FILL_REVIEW = "FFE8D6"
FILL_HEADER = "0A1F3D"  # primary.900: the navy of the web app
FILL_HELPER_HEADER = "33415A"  # the helper columns' header is a step quieter
FILL_ZEBRA = "F7F9FC"
BORDER_COLOUR = "DCE4EF"  # surface.border
FONT_NAME = "Arial"
HEADER_HEIGHT = 34
ROW_HEIGHT = 24
# Readable widths per column (characters), in COLUMNS order; the reason column wraps.
WIDTHS = (14, 16, 16, 22, 16, 24, 20, 12, 12, 30, 18, 26, 60, 12)
WRAP_COLUMNS = frozenset({13})

DATE_FORMAT = "DD/MM/YYYY"
SOURCE_DATE_FORMAT = "%d/%m/%Y"
TEXT_FORMAT = "@"
# Identifiers must stay text: Excel would drop the zero of 011223344 and could mangle X132767.
TEXT_KEYS: frozenset[FieldKey | str] = frozenset(
    {FieldKey.national_id, CENTRALS_KEY, FieldKey.commitment_number}
)

# Structural: bounds on content-derived widths so one long reason cannot make a column unusable.
MIN_WIDTH = 10
MAX_WIDTH = 80
WIDTH_PADDING = 3
COMMENT_AUTHOR = "Safe Line"
COMMENT_SIZE = (320, 110)  # width, height in points

AUDIT_HEADERS = ("זמן", "שורה", "שדה", "ערך קודם", "ערך חדש", "פעולה", "מבצע")
LEGEND = (
    (FILL_VERIFIED, "ירוק - אומת פעמיים או בברקוד"),
    (FILL_SINGLE, "ירוק בהיר - זוהה בקריאה אחת ועבר את כל הבדיקות"),
    (FILL_REVIEW, "כתום - נדרשת בדיקה"),
)
SUMMARY_LABELS = (
    ("page_count", "דפים"),
    ("row_count", "שורות"),
    ("fields_total", "שדות"),
    ("fields_closed_local", "שדות שנסגרו מקומית"),
    ("fields_escalated", "שדות שהועלו לשלב הבא"),
    ("fields_flagged", "שדות לבדיקה"),
    ("tier1_calls", "קריאות שכבה 1"),
    ("tier2_calls", "קריאות שכבה 2"),
    ("local_close_rate", "אחוז סגירה מקומית"),
    ("engine_version", "גרסת מנוע"),
)


def _fill_for(level: ConfidenceLevel | None) -> PatternFill | None:
    if level is None:
        return None
    if level in (ConfidenceLevel.A, ConfidenceLevel.B):
        colour = FILL_VERIFIED
    elif level is ConfidenceLevel.C:
        colour = FILL_SINGLE
    else:
        colour = FILL_REVIEW
    return PatternFill("solid", start_color=colour, end_color=colour)


def _set_text(cell: object, text: str) -> None:
    """Write a string that Excel must never reinterpret, even when it starts with '='."""
    cell.value = text  # type: ignore[attr-defined]
    cell.data_type = "s"  # type: ignore[attr-defined]


_THIN = Side(style="thin", color=BORDER_COLOUR)
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _write_value(ws: Worksheet, row_idx: int, col_idx: int, key: FieldKey | str, c: Cell) -> None:
    target = ws.cell(row=row_idx, column=col_idx)
    target.border = _BORDER
    target.font = Font(name=FONT_NAME, size=11, bold=col_idx <= 7 and bool(c.text))
    parsed = _parse_date(c.text) if key is FieldKey.receipt_date else None
    if parsed is not None:
        target.value = parsed
        target.number_format = DATE_FORMAT
    else:
        _set_text(target, c.text)
        if key in TEXT_KEYS:
            target.number_format = TEXT_FORMAT
    fill = _fill_for(c.level)
    if fill is not None:
        target.fill = fill
    elif row_idx % 2 == 1:
        target.fill = PatternFill("solid", start_color=FILL_ZEBRA, end_color=FILL_ZEBRA)
    if c.comment:
        comment = Comment(c.comment, COMMENT_AUTHOR)
        comment.width, comment.height = COMMENT_SIZE
        target.comment = comment
    target.alignment = Alignment(
        horizontal="right", vertical="center", wrap_text=col_idx in WRAP_COLUMNS
    )


def _parse_date(text: str) -> datetime | None:
    """A real date when the text is DD/MM/YYYY; otherwise None and the text is kept as is."""
    try:
        return datetime.strptime(text, SOURCE_DATE_FORMAT)
    except ValueError:
        return None


def _style_header(ws: Worksheet, headers: tuple[str, ...]) -> None:
    for idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=idx, value=header)
        colour = FILL_HEADER if idx <= 7 or len(headers) < 8 else FILL_HELPER_HEADER
        cell.font = Font(name=FONT_NAME, bold=True, size=11, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color=colour, end_color=colour)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
    ws.row_dimensions[1].height = HEADER_HEIGHT
    ws.freeze_panes = "A2"
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.zoomScale = 110
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _fit_widths(ws: Worksheet, column_count: int) -> None:
    """Width from the longest line in each column, so nothing is truncated (up to the cap)."""
    for idx in range(1, column_count + 1):
        longest = 0
        for row in ws.iter_rows(min_col=idx, max_col=idx):
            for cell in row:
                if cell.value is not None:
                    shown = DATE_FORMAT if isinstance(cell.value, datetime) else str(cell.value)
                    longest = max(longest, len(shown))
        width = min(MAX_WIDTH, max(MIN_WIDTH, longest + WIDTH_PADDING))
        ws.column_dimensions[get_column_letter(idx)].width = width


def _fill_table(ws: Worksheet, flat_rows: list[FlatRow]) -> None:
    _style_header(ws, HEADERS)
    for r_idx, flat in enumerate(flat_rows, start=2):
        for c_idx, ((_, key), cell) in enumerate(zip(COLUMNS, flat.cells, strict=True), start=1):
            _write_value(ws, r_idx, c_idx, key, cell)
        ws.row_dimensions[r_idx].height = ROW_HEIGHT
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(1, len(flat_rows) + 1)}"
    for idx, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def _fill_audit(ws: Worksheet) -> None:
    """Headers only: ExporterPort carries no audit log; the column set is fixed for later."""
    _style_header(ws, AUDIT_HEADERS)
    _fit_widths(ws, len(AUDIT_HEADERS))


def _section(ws: Worksheet, row: int, title: str) -> None:
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = Font(name=FONT_NAME, bold=True, size=12, color="FFFFFF")
    for col in (1, 2):
        ws.cell(row=row, column=col).fill = PatternFill(
            "solid", start_color=FILL_HEADER, end_color=FILL_HEADER
        )
    ws.row_dimensions[row].height = ROW_HEIGHT + 4


def _pair(ws: Worksheet, row: int, label: str, value: object, fill: str | None = None) -> None:
    left = ws.cell(row=row, column=1, value=label)
    right = ws.cell(row=row, column=2, value=value)
    for cell in (left, right):
        cell.font = Font(name=FONT_NAME, size=11)
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="right", vertical="center")
    if fill is not None:
        left.fill = PatternFill("solid", start_color=fill, end_color=fill)
    ws.row_dimensions[row].height = ROW_HEIGHT


def _fill_summary(ws: Worksheet, report: RunReport) -> None:
    """Title, colour legend and RunReport counters. Counters only - no field values."""
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    title = ws.cell(row=1, column=1, value="שורה בטוחה · סיכום ריצה")
    title.font = Font(name=FONT_NAME, bold=True, size=16, color=FILL_HEADER)
    ws.row_dimensions[1].height = 30
    _section(ws, 3, "מקרא צבעים")
    row = 4
    for colour, text in LEGEND:
        _pair(ws, row, text, None, colour)
        row += 1
    row += 1
    _section(ws, row, "נתוני ריצה")
    row += 1
    for attr, label in SUMMARY_LABELS:
        _pair(ws, row, label, getattr(report, attr))
        row += 1
    for level, count in sorted(report.level_histogram.items()):
        _pair(ws, row, f"רמה {level}", count)
        row += 1
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 22


class XlsxExporter:
    fmt = "xlsx"
    media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def render(self, report: RunReport, rows: list[CommitmentRow]) -> bytes:
        """Four right-to-left sheets built in memory. An empty job yields a valid workbook."""
        flat_rows = [flatten(row) for row in rows]
        wb = Workbook()
        main = wb.active
        assert main is not None  # a new Workbook always has an active sheet
        main.title = SHEET_MAIN
        _fill_table(main, flat_rows)
        _fill_table(wb.create_sheet(SHEET_REVIEW), [f for f in flat_rows if f.flagged])
        _fill_audit(wb.create_sheet(SHEET_AUDIT))
        _fill_summary(wb.create_sheet(SHEET_SUMMARY), report)
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()
