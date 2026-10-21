from __future__ import annotations

from datetime import datetime
import io

from openpyxl import Workbook, load_workbook
import pytest

from backend.adapters.export.table import HEADERS
from backend.adapters.export.xlsx import (
    FILL_REVIEW,
    FILL_SINGLE,
    FILL_VERIFIED,
    SHEET_AUDIT,
    SHEET_MAIN,
    SHEET_REVIEW,
    SHEET_SUMMARY,
    XlsxExporter,
)
from backend.core.models import CommitmentRow, ConfidenceLevel, FieldKey
from backend.tests.unit.export_helpers import make_field, make_report, make_row

COL_ID, COL_CENTRAL, COL_DATE, COL_COMMIT = 3, 4, 5, 7


def _open(rows: list[CommitmentRow]) -> Workbook:
    return load_workbook(io.BytesIO(XlsxExporter().render(make_report(), rows)))


def _with_commitment(
    row: CommitmentRow, level: ConfidenceLevel, value: str | None
) -> CommitmentRow:
    fields = dict(row.fields)
    fields[FieldKey.commitment_number] = make_field(FieldKey.commitment_number, value, level)
    return row.model_copy(update={"fields": fields})


def test_national_id_keeps_leading_zero() -> None:
    cell = _open([make_row()])[SHEET_MAIN].cell(row=2, column=COL_ID)
    assert cell.value == "011223344"
    assert cell.number_format == "@"


def test_commitment_with_letter_survives() -> None:
    cell = _open([make_row()])[SHEET_MAIN].cell(row=2, column=COL_COMMIT)
    assert cell.value == "X132767"
    assert cell.number_format == "@"


def test_receipt_date_is_a_real_date() -> None:
    cell = _open([make_row()])[SHEET_MAIN].cell(row=2, column=COL_DATE)
    assert isinstance(cell.value, datetime)
    assert (cell.value.year, cell.value.month, cell.value.day) == (2026, 3, 5)
    assert cell.number_format == "DD/MM/YYYY"


def test_unparseable_date_is_kept_as_text() -> None:
    cell = _open([make_row(receipt_date="5 במרץ")])[SHEET_MAIN].cell(row=2, column=COL_DATE)
    assert cell.value == "5 במרץ"


def test_multiple_centrals_joined_with_semicolon() -> None:
    cell = _open([make_row(centrals=("111", "222"))])[SHEET_MAIN].cell(row=2, column=COL_CENTRAL)
    assert cell.value == "111; 222"


@pytest.mark.parametrize(
    ("level", "colour"),
    [
        (ConfidenceLevel.A, FILL_VERIFIED),
        (ConfidenceLevel.B, FILL_VERIFIED),
        (ConfidenceLevel.C, FILL_SINGLE),
        (ConfidenceLevel.D, FILL_REVIEW),
    ],
)
def test_cell_fill_matches_confidence_level(level: ConfidenceLevel, colour: str) -> None:
    row = _with_commitment(make_row(), level, "X132767")
    cell = _open([row])[SHEET_MAIN].cell(row=2, column=COL_COMMIT)
    assert cell.fill.start_color.rgb.endswith(colour)


def test_orange_cell_has_comment_with_reason() -> None:
    row = _with_commitment(make_row(), ConfidenceLevel.D, None)
    cell = _open([row])[SHEET_MAIN].cell(row=2, column=COL_COMMIT)
    assert cell.value in (None, "")
    assert cell.comment is not None
    assert "לא נמצא בטופס" in cell.comment.text


def test_green_cell_has_no_comment() -> None:
    cell = _open([make_row()])[SHEET_MAIN].cell(row=2, column=COL_COMMIT)
    assert cell.comment is None


def test_sheets_are_rtl() -> None:
    wb = _open([make_row()])
    assert all(ws.sheet_view.rightToLeft for ws in wb.worksheets)


def test_four_sheets_present_with_hebrew_names() -> None:
    assert _open([]).sheetnames == [SHEET_MAIN, SHEET_REVIEW, SHEET_AUDIT, SHEET_SUMMARY]


def test_review_sheet_contains_only_flagged_rows() -> None:
    ok = make_row("ok")
    bad = _with_commitment(make_row("bad"), ConfidenceLevel.D, None)
    review = _open([ok, bad])[SHEET_REVIEW]
    assert review.max_row == 2
    assert review.cell(row=2, column=COL_ID).value == "011223344"
    assert review.cell(row=2, column=12).value == "מספר התחייבות"


def test_summary_sheet_has_no_field_values() -> None:
    wb = _open([make_row()])
    cells = {str(c.value) for row in wb[SHEET_SUMMARY].iter_rows() for c in row if c.value}
    assert not cells & {"011223344", "X132767", "דנה", "כהן", "123456"}


def test_column_order_matches_guidelines() -> None:
    header = tuple(c.value for c in _open([])[SHEET_MAIN][1])
    assert header == HEADERS
    assert header[:7] == (
        "שם פרטי",
        "שם משפחה",
        "מספר תעודת זהות",
        "מספר מרכזי",
        "תאריך קבלת הדגימה",
        "גורם שולח",
        "מספר התחייבות",
    )


def test_formula_like_text_stays_text() -> None:
    row = _with_commitment(make_row(), ConfidenceLevel.C, "=1+1")
    cell = _open([row])[SHEET_MAIN].cell(row=2, column=COL_COMMIT)
    assert cell.value == "=1+1"
    assert cell.data_type == "s"


def test_header_is_frozen_with_autofilter() -> None:
    ws = _open([make_row()])[SHEET_MAIN]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:N2"
